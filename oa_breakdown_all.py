"""
Full Open Access breakdown for ALL institutions (8 intl orgs + 50 US)
Scope: 2018–2026
Generates: 3 charts + updated spreadsheet
"""
import requests, socket, sys, os, time
import matplotlib.pyplot as plt
import matplotlib.ticker as ticker
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import urllib3.util.connection as u
u.allowed_gai_family = lambda: socket.AF_INET
sys.stdout.reconfigure(encoding='utf-8')

API_KEY = "eN9RGcYK6NQuQJHOJ2XBDX"
BASE = "https://api.openalex.org"
YEAR_FILTER = "from_publication_date:2018-01-01"
YEARS = list(range(2018, 2027))
OUTPUT_DIR = "C:/Users/talay/Documents/openalex_project/charts_2018"
XLSX_PATH = "C:/Users/talay/Documents/openalex_project/OpenAlex_OA_Breakdown_All.xlsx"
os.makedirs(OUTPUT_DIR, exist_ok=True)

OA_FLAVORS = ["gold", "green", "hybrid", "bronze", "diamond", "closed"]
OA_COLORS = {
    "gold": "#F59E0B", "green": "#10B981", "hybrid": "#8B5CF6",
    "bronze": "#F97316", "diamond": "#06B6D4", "closed": "#9CA3AF",
}
OA_LABELS = {
    "gold": "Gold", "green": "Green", "hybrid": "Hybrid",
    "bronze": "Bronze", "diamond": "Diamond", "closed": "Closed",
}

plt.rcParams.update({
    'figure.facecolor': '#f8f9fa', 'axes.facecolor': '#ffffff',
    'axes.grid': True, 'grid.alpha': 0.3, 'font.size': 10,
    'axes.titlesize': 13, 'axes.titleweight': 'bold',
    'figure.titlesize': 15, 'figure.titleweight': 'bold',
})

H_FONT = Font(name='Arial', bold=True, color='FFFFFF', size=11)
H_FILL = PatternFill('solid', fgColor='2563EB')
D_FONT = Font(name='Arial', size=10)
B_FONT = Font(name='Arial', bold=True, size=10)
L_FONT = Font(name='Arial', size=10, color='2563EB', underline='single')
NUM = '#,##0'; DEC = '0.0'; PCT = '0.0%'
BDR = Border(bottom=Side(style='thin', color='D1D5DB'), right=Side(style='thin', color='D1D5DB'))
CTR = Alignment(horizontal='center', vertical='center')
LFT = Alignment(horizontal='left', vertical='center')
WRP = Alignment(horizontal='center', vertical='center', wrap_text=True)


def api(endpoint, params=None, retries=3):
    if params is None: params = {}
    params["api_key"] = API_KEY
    for attempt in range(retries):
        try:
            r = requests.get(f"{BASE}/{endpoint}", params=params, timeout=30)
            r.raise_for_status(); return r.json()
        except Exception as e:
            if attempt < retries - 1: time.sleep(1)
            else: print(f"    [!] {e}"); return {}


def shorten(name, mx=28):
    reps = [
        ("University of California, ", "UC "), ("University of California ", "UC "),
        ("The University of Texas at ", "UT "), ("University of North Carolina at ", "UNC "),
        ("University of Illinois Urbana-Champaign", "UIUC"),
        ("University of Illinois Chicago", "UIC"),
        ("Rutgers, The State University of New Jersey", "Rutgers"),
        ("Purdue University West Lafayette", "Purdue"),
        ("Indiana University Bloomington", "Indiana U"),
        ("Indiana University – Purdue University Indianapolis", "IUPUI"),
        ("Washington University in St. Louis", "WashU St. Louis"),
        ("Massachusetts Institute of Technology", "MIT"),
        ("Massachusetts General Hospital", "Mass General"),
        ("The Ohio State University", "Ohio State"),
        ("Pennsylvania State University", "Penn State"),
        ("Michigan State University", "Michigan State"),
        ("George Washington University", "George Washington U"),
        ("Texas A&M University", "Texas A&M"),
        ("International Monetary Fund", "IMF"),
        ("Food and Agriculture Organization", "FAO"),
        ("International Labour Organization", "ILO"),
        ("International Telecommunication Union", "ITU"),
        ("World Intellectual Property Organization", "WIPO"),
        ("United Nations Development Programme", "UNDP"),
        ("United Nations", "UN"), ("World Bank", "World Bank"),
        ("University of ", "U "), ("University", "U"),
    ]
    for old, new in reps:
        if old in name: name = name.replace(old, new); break
    return name[:mx-2] + ".." if len(name) > mx else name


def save_fig(fig, fn):
    p = f"{OUTPUT_DIR}/{fn}"
    fig.savefig(p, dpi=150, bbox_inches='tight', facecolor=fig.get_facecolor())
    plt.close(fig); print(f"    Saved: {p}")


def style_header(ws, row, mc):
    for c in range(1, mc+1):
        cell = ws.cell(row=row, column=c)
        cell.font = H_FONT; cell.fill = H_FILL; cell.alignment = WRP; cell.border = BDR


def auto_width(ws, mn=10, mx=38):
    for col_cells in ws.columns:
        cl = get_column_letter(col_cells[0].column)
        w = max((len(str(c.value or '')) for c in col_cells), default=mn)
        ws.column_dimensions[cl].width = min(max(w+2, mn), mx)


# ==================================================================
print("\n" + "=" * 65)
print("  OA BREAKDOWN — ALL INSTITUTIONS — 2018 ONWARDS")
print("=" * 65)

# ── 1. Identify international organizations ──
INTL_ORGS = [
    "United Nations", "World Bank", "International Monetary Fund",
    "Food and Agriculture Organization", "International Labour Organization",
    "International Telecommunication Union", "World Intellectual Property Organization",
    "United Nations Development Programme",
]

print("\n  [1/4] Fetching intl org profiles...")
intl = []
for org_name in INTL_ORGS:
    print(f"    -> {org_name}")
    data = api("institutions", {"search": org_name, "per_page": 5})
    inst = data.get("results", [None])[0]
    for r in data.get("results", []):
        if r.get("type") in ("government", "nonprofit", "facility", "other"):
            inst = r; break
    if not inst: continue
    intl.append({
        "id": inst["id"].split("/")[-1],
        "name": inst["display_name"],
        "group": "intl",
    })

# ── 2. Identify top 50 US SS institutions ──
print("\n  [2/4] Fetching top 50 US SS institution list...")
ss_raw = api("works", {
    "filter": f"concepts.id:C36289849,authorships.institutions.country_code:US,"
              f"authorships.institutions.type:education,{YEAR_FILTER}",
    "group_by": "authorships.institutions.id", "per_page": 50,
})
ss_rankings = [(g["key_display_name"], g["key"], g["count"]) for g in ss_raw.get("group_by", [])]

us = []
for name, key, ss_count in ss_rankings:
    inst_id = key.split("/")[-1]
    us.append({
        "id": inst_id,
        "name": name,
        "group": "us",
        "ss_works": ss_count,
    })

# ── 3. Fetch OA breakdowns for ALL institutions ──
all_insts = intl + us
print(f"\n  [3/4] Fetching OA breakdowns for {len(all_insts)} institutions...")

oa_data = []
for i, inst in enumerate(all_insts, 1):
    label = f"[{i}/{len(all_insts)}]"
    print(f"    {label} {inst['name']}")
    data = api("works", {
        "filter": f"authorships.institutions.id:{inst['id']},{YEAR_FILTER}",
        "group_by": "open_access.oa_status"
    })
    counts = {g["key"]: g["count"] for g in data.get("group_by", [])}
    total = sum(counts.values())
    pcts = {}
    for flavor in OA_FLAVORS:
        c = counts.get(flavor, 0)
        pcts[flavor] = round(100 * c / total, 1) if total else 0
        pcts[f"{flavor}_n"] = c
    pcts["total"] = total
    pcts["open_total"] = total - counts.get("closed", 0)
    pcts["open_pct"] = round(100 * pcts["open_total"] / total, 1) if total else 0
    oa_data.append({**inst, **pcts})
    time.sleep(0.05)

# ── 4. Charts & Spreadsheet ──
print(f"\n  [4/4] Generating charts & spreadsheet...")

intl_oa = [d for d in oa_data if d["group"] == "intl"]
us_oa = [d for d in oa_data if d["group"] == "us"]

# --- Chart A: Intl Orgs OA stacked bar ---
fig, ax = plt.subplots(figsize=(12, 6))
fig.suptitle("International Organizations — Open Access Breakdown 2018–2026")
names = [shorten(d["name"], 22) for d in intl_oa]
y_pos = np.arange(len(names))
left = np.zeros(len(names))
for flavor in OA_FLAVORS:
    vals = [d[flavor] for d in intl_oa]
    bars = ax.barh(y_pos, vals, left=left, color=OA_COLORS[flavor],
                   label=OA_LABELS[flavor], height=0.65, edgecolor='white', linewidth=0.5)
    for j, (v, l) in enumerate(zip(vals, left)):
        if v >= 4:
            ax.text(l + v/2, j, f"{v:.0f}%", ha='center', va='center',
                    fontsize=8, fontweight='bold', color='white' if flavor == 'closed' else '#1a1a1a')
    left += vals
ax.set_yticks(y_pos); ax.set_yticklabels(names, fontsize=10)
ax.set_xlabel("% of Publications (2018+)")
ax.set_xlim(0, 100)
ax.legend(loc='lower right', fontsize=9, ncol=3, framealpha=0.9)
ax.invert_yaxis()
fig.tight_layout(); save_fig(fig, "10_intl_oa_breakdown_2018.png")

# --- Chart B: Top 50 US OA stacked bar ---
fig, ax = plt.subplots(figsize=(14, 22))
fig.suptitle("Top 50 US Social Science — Open Access Breakdown 2018–2026", y=0.995)
names = [shorten(d["name"], 26) for d in us_oa]
y_pos = np.arange(len(names))
left = np.zeros(len(names))
for flavor in OA_FLAVORS:
    vals = [d[flavor] for d in us_oa]
    bars = ax.barh(y_pos, vals, left=left, color=OA_COLORS[flavor],
                   label=OA_LABELS[flavor], height=0.7, edgecolor='white', linewidth=0.3)
    for j, (v, l) in enumerate(zip(vals, left)):
        if v >= 5:
            ax.text(l + v/2, j, f"{v:.0f}%", ha='center', va='center',
                    fontsize=7, fontweight='bold', color='white' if flavor == 'closed' else '#1a1a1a')
    left += vals
ax.set_yticks(y_pos); ax.set_yticklabels(names, fontsize=8)
ax.set_xlabel("% of Publications (2018+)")
ax.set_xlim(0, 100)
ax.legend(loc='lower right', fontsize=9, ncol=3, framealpha=0.9)
ax.invert_yaxis()
fig.tight_layout(); save_fig(fig, "11_top50_us_oa_breakdown_2018.png")

# --- Chart C: Open Access Rate comparison (sorted by OA%) ---
all_sorted = sorted(oa_data, key=lambda x: x["open_pct"], reverse=True)
fig, ax = plt.subplots(figsize=(14, 24))
fig.suptitle("All Institutions — Open Access Rate 2018–2026", y=0.997)
names = [shorten(d["name"], 26) for d in all_sorted]
y_pos = np.arange(len(names))
colors = ['#2563eb' if d["group"] == "intl" else '#16a34a' for d in all_sorted]
oa_pcts = [d["open_pct"] for d in all_sorted]
bars = ax.barh(y_pos, oa_pcts, color=colors, height=0.7, edgecolor='white', linewidth=0.3)
for j, (v, bar) in enumerate(zip(oa_pcts, bars)):
    ax.text(v + 0.5, j, f"{v:.1f}%", ha='left', va='center', fontsize=7.5)
ax.set_yticks(y_pos); ax.set_yticklabels(names, fontsize=8)
ax.set_xlabel("Open Access Rate (%)")
ax.set_xlim(0, 105)
# Custom legend
from matplotlib.patches import Patch
legend_elements = [Patch(facecolor='#2563eb', label='Intl Organizations'),
                   Patch(facecolor='#16a34a', label='US Academic Institutions')]
ax.legend(handles=legend_elements, loc='lower right', fontsize=10, framealpha=0.9)
ax.invert_yaxis()
fig.tight_layout(); save_fig(fig, "12_all_oa_rate_2018.png")

# --- Spreadsheet ---
print("\n    Building spreadsheet...")
wb = Workbook()

# Sheet 1: Intl Orgs OA
ws1 = wb.active
ws1.title = "Intl Orgs OA"
headers1 = ["Organization", "Total Works", "Open Access Works", "OA Rate %",
            "Gold %", "Green %", "Hybrid %", "Bronze %", "Diamond %", "Closed %",
            "Gold", "Green", "Hybrid", "Bronze", "Diamond", "Closed", "OpenAlex Link"]
ws1.append(headers1); style_header(ws1, 1, len(headers1))
for d in intl_oa:
    r = ws1.max_row + 1
    ws1.cell(r, 1, d["name"])
    ws1.cell(r, 2, d["total"]); ws1.cell(r, 2).number_format = NUM
    ws1.cell(r, 3, d["open_total"]); ws1.cell(r, 3).number_format = NUM
    ws1.cell(r, 4, d["open_pct"]/100); ws1.cell(r, 4).number_format = PCT
    for ci, fl in enumerate(OA_FLAVORS):
        ws1.cell(r, 5+ci, d[fl]/100); ws1.cell(r, 5+ci).number_format = PCT
        ws1.cell(r, 11+ci, d[f"{fl}_n"]); ws1.cell(r, 11+ci).number_format = NUM
    link = f"https://api.openalex.org/works?filter=authorships.institutions.id:{d['id']},from_publication_date:2018-01-01&group_by=open_access.oa_status&api_key={API_KEY}"
    ws1.cell(r, 17, "API Link"); ws1.cell(r, 17).hyperlink = link
    ws1.cell(r, 17).font = L_FONT
    for c in range(1, len(headers1)+1):
        ws1.cell(r, c).border = BDR
        ws1.cell(r, c).alignment = CTR if c > 1 else LFT
auto_width(ws1)

# Sheet 2: Top 50 US OA
ws2 = wb.create_sheet("Top 50 US SS OA")
headers2 = ["Rank", "Institution", "SS Works (2018+)", "Total Works", "Open Access Works", "OA Rate %",
            "Gold %", "Green %", "Hybrid %", "Bronze %", "Diamond %", "Closed %",
            "Gold", "Green", "Hybrid", "Bronze", "Diamond", "Closed", "OpenAlex Link"]
ws2.append(headers2); style_header(ws2, 1, len(headers2))
for rank, d in enumerate(us_oa, 1):
    r = ws2.max_row + 1
    ws2.cell(r, 1, rank)
    ws2.cell(r, 2, d["name"])
    ws2.cell(r, 3, d.get("ss_works", 0)); ws2.cell(r, 3).number_format = NUM
    ws2.cell(r, 4, d["total"]); ws2.cell(r, 4).number_format = NUM
    ws2.cell(r, 5, d["open_total"]); ws2.cell(r, 5).number_format = NUM
    ws2.cell(r, 6, d["open_pct"]/100); ws2.cell(r, 6).number_format = PCT
    for ci, fl in enumerate(OA_FLAVORS):
        ws2.cell(r, 7+ci, d[fl]/100); ws2.cell(r, 7+ci).number_format = PCT
        ws2.cell(r, 13+ci, d[f"{fl}_n"]); ws2.cell(r, 13+ci).number_format = NUM
    link = f"https://api.openalex.org/works?filter=authorships.institutions.id:{d['id']},from_publication_date:2018-01-01&group_by=open_access.oa_status&api_key={API_KEY}"
    ws2.cell(r, 19, "API Link"); ws2.cell(r, 19).hyperlink = link
    ws2.cell(r, 19).font = L_FONT
    for c in range(1, len(headers2)+1):
        ws2.cell(r, c).border = BDR
        ws2.cell(r, c).alignment = CTR if c > 1 else LFT
auto_width(ws2)

# Sheet 3: Combined ranking by OA rate
ws3 = wb.create_sheet("OA Rate Ranking")
headers3 = ["Rank", "Institution", "Type", "Total Works", "OA Rate %",
            "Gold %", "Green %", "Hybrid %", "Bronze %", "Diamond %", "Closed %"]
ws3.append(headers3); style_header(ws3, 1, len(headers3))
for rank, d in enumerate(all_sorted, 1):
    r = ws3.max_row + 1
    ws3.cell(r, 1, rank)
    ws3.cell(r, 2, d["name"])
    ws3.cell(r, 3, "Intl Org" if d["group"] == "intl" else "US Academic")
    ws3.cell(r, 4, d["total"]); ws3.cell(r, 4).number_format = NUM
    ws3.cell(r, 5, d["open_pct"]/100); ws3.cell(r, 5).number_format = PCT
    for ci, fl in enumerate(OA_FLAVORS):
        ws3.cell(r, 6+ci, d[fl]/100); ws3.cell(r, 6+ci).number_format = PCT
    for c in range(1, len(headers3)+1):
        ws3.cell(r, c).border = BDR
        ws3.cell(r, c).alignment = CTR if c > 1 else LFT
auto_width(ws3)

wb.save(XLSX_PATH)
print(f"    Saved: {XLSX_PATH}")

print("\n" + "=" * 65)
print("  DONE — 3 charts + spreadsheet generated")
print(f"  Charts: {OUTPUT_DIR}")
print(f"  Spreadsheet: {XLSX_PATH}")
print("=" * 65)
